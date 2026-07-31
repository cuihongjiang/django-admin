# -*- coding: utf-8 -*-
# @Time    : 2024/12/7 00:24
# @Author  : 崔宏江
# @FileName: js_crud.py
# @Software: VSCode
"""
Excel 导入导出工具（DRF 版）
ninja 时代的 create/update/retrieve 等通用 CRUD 已由
DRF Serializer + CoreModelViewSet 取代，此处仅保留 Excel 能力
"""
import io
from datetime import datetime

import openpyxl
from django.http import FileResponse
from openpyxl import load_workbook


def export_excel(queryset, serializer_class, export_fields, model):
    """
    导出查询集为 Excel 文件（内存生成，不落盘）

    参数:
    - queryset: 要导出的查询集
    - serializer_class: DRF 序列化器类，用于取字段值
    - export_fields: 要导出的字段名列表
    - model: 模型类，用于取字段 help_text 作为表头
    """
    # 字段名 -> 表头显示名（help_text）
    title_dict = {}
    for field in export_fields:
        field_obj = model._meta.get_field(field)
        title_dict[field] = str(field_obj.help_text or field)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(title_dict.values()))
    for instance in queryset:
        row = serializer_class(instance).data
        ws.append([_cell_value(row.get(field)) for field in title_dict])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    file_name = datetime.now().strftime('%Y%m%d%H%M%S%f') + '.xlsx'
    return FileResponse(buffer, as_attachment=True, filename=file_name)


def import_excel(file_obj, model, import_fields, request):
    """
    从上传的 Excel 文件导入数据

    参数:
    - file_obj: 上传的 Excel 文件对象（request.FILES 中取得）
    - model: 目标模型类
    - import_fields: 允许导入的字段名列表
    - request: 请求对象，用于填充审计字段

    返回值:
    - 成功导入的记录条数
    """
    # 表头显示名（help_text）-> 字段名
    title_dict = {}
    for field in import_fields:
        field_obj = model._meta.get_field(field)
        title_dict[str(field_obj.help_text or field)] = field

    wb = load_workbook(file_obj)
    ws = wb.active
    rows = ws.values
    header = next(rows, None)
    if header is None:
        return 0

    user = request.user
    count = 0
    for row in rows:
        data = {}
        for index, cell in enumerate(row):
            field = title_dict.get(header[index]) if index < len(header) else None
            if field is not None and cell is not None:
                data[field] = cell
        if not data:
            continue
        model.objects.create(
            **data,
            creator=user if getattr(user, 'is_authenticated', False) else None,
            modifier=getattr(user, 'username', None),
            belong_dept=getattr(user, 'dept_id', None),
        )
        count += 1
    return count


def _cell_value(value):
    """Excel 单元格仅支持标量，列表（如多对多 id）转为逗号分隔字符串"""
    if isinstance(value, (list, tuple)):
        return ','.join(str(v) for v in value)
    if isinstance(value, (dict, set)):
        return str(value)
    return value
